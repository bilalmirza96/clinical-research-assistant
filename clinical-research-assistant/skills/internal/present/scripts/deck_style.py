"""
CRA conference-deck design system.

Constants were measured off a real 16:9 clinical-research deck (ITSOS 2026), so slides
built with this module are visually consistent with a hand-built academic deck. Native
PowerPoint shapes, not images, wherever the content is text or simple geometry -- native
shapes stay editable by the author and re-render crisply at any projector resolution.
Charts that need a real plotting engine (Kaplan-Meier, distributions) are rendered with
matplotlib and placed full-bleed; see references/slide-patterns.md.

Slide is 20.00 x 11.25 in (16:9). All geometry in inches.

PALETTE NOTE: NHW/NHB are the project race convention (BLUE = non-Hispanic White,
RED = non-Hispanic Black; RED is RESERVED for NHB). Registry/category series use
NCDB_BLUE / SEER_ORANGE or ALT -- never red while a race series is on the same slide.
"""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import copy

# ---- palette (measured) -----------------------------------------------------
NHW      = RGBColor(0x12, 0x30, 0x57)   # non-Hispanic White  -- navy
NHB      = RGBColor(0x80, 0x18, 0x19)   # non-Hispanic Black  -- dark red
ALT      = RGBColor(0x51, 0x80, 0xBF)   # secondary series (registry/category)
NCDB_BLUE   = RGBColor(0x00, 0x7A, 0xD1)  # registry: NCDB
SEER_ORANGE = RGBColor(0xD1, 0x56, 0x00)  # registry: SEER
INK      = RGBColor(0x22, 0x33, 0x3F)   # headers, category labels
MUTED    = RGBColor(0x4B, 0x55, 0x63)   # axis labels, tick labels
FOOT     = RGBColor(0x5B, 0x66, 0x75)   # source line, slide number
GRID     = RGBColor(0xE9, 0xEC, 0xF0)   # gridlines
RULE     = RGBColor(0xE1, 0xE6, 0xED)   # title / footer rules
AXIS     = RGBColor(0x6B, 0x72, 0x80)   # baseline
PALE     = RGBColor(0xC5, 0xCD, 0xD8)   # 4th stacked segment
BLACK    = RGBColor(0x00, 0x00, 0x00)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Times New Roman"

# ---- type scale (measured) --------------------------------------------------
SZ_TITLE=43.5; SZ_PANEL=22.5; SZ_VALUE=21.0; SZ_CAT=19.5
SZ_AXIS=18.0;  SZ_LEGEND=18.0; SZ_EST=22.5; SZ_SUB=16.5; SZ_SRC=16.5

# ---- canonical geometry (measured) ------------------------------------------
TITLE=(0.94,0.73,19.94,0.71); TITLE_RULE=(0.94,1.62,18.12,0.01)
PANEL_A_LET=(0.90,2.11,0.50,0.46); PANEL_A_HDR=(1.52,2.11,8.25)
PANEL_B_LET=(10.43,2.11,0.50,0.46); PANEL_B_HDR=(11.05,2.11,8.25)
FOOT_RULE=(0.94,10.29,18.12,0.01)
LEG_Y=10.50; LEG_TXT_Y=10.45; SRC=(6.56,10.45,8.91,0.31); PGNUM=(18.83,10.45,0.31,0.31)

# left bar panel
BAR_X0=2.29; BAR_W=6.25; BAR_BASE=8.10; BAR_TOP=3.94; BARW=0.70
# right forest panel
FOR_X0=12.05; FOR_X1=15.44; FOR_LBL=16.09; FOR_TOP=4.21; FOR_DY=1.25
# stacked composition
STK_X0=4.48; STK_W=13.54; STK_H=1.15


def _tx(slide,l,t,w,h,text,size,color,bold=False,align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP,wrap=True,italic=False):
    b=slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    tf=b.text_frame; tf.word_wrap=wrap
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text
    f=r.font; f.name=FONT; f.size=Pt(size); f.bold=bold; f.italic=italic
    f.color.rgb=color
    return b


def _rect(slide,l,t,w,h,color,shape=MSO_SHAPE.RECTANGLE,adj=None):
    s=slide.shapes.add_shape(shape,Inches(l),Inches(t),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=color; s.line.fill.background()
    s.shadow.inherit=False
    if adj is not None:
        try: s.adjustments[0]=adj
        except Exception: pass
    return s


def frame(slide,title,source,pagenum,legend=None):
    """Title, rules, source line, page number and optional race legend."""
    _tx(slide,*TITLE,title,SZ_TITLE,BLACK,bold=True)
    _rect(slide,*TITLE_RULE,RULE)
    _rect(slide,*FOOT_RULE,RULE)
    if legend:
        # swatch/label x-positions measured off the deck's own slide 20 so the legend
        # never runs into the source line that starts at x=6.56
        SLOT=[(0.94,1.31,2.15),(3.76,4.13,2.13)]
        for i,(lab,col) in enumerate(legend[:2]):
            sx,tx,tw=SLOT[i]
            _rect(slide,sx,LEG_Y,0.27,0.17,col)
            _tx(slide,tx,LEG_TXT_Y,tw,0.32,lab,SZ_LEGEND,INK,wrap=False)
    _tx(slide,*SRC,source,SZ_SRC,FOOT)
    _tx(slide,*PGNUM,str(pagenum),SZ_SRC,FOOT,align=PP_ALIGN.RIGHT)


def panel_header(slide,letter,text,side="A",height=1.26,width=None):
    lt = PANEL_A_LET if side=="A" else PANEL_B_LET
    hd = PANEL_A_HDR if side=="A" else PANEL_B_HDR
    _tx(slide,*lt,letter,SZ_PANEL,INK,bold=True)
    _tx(slide,hd[0],hd[1],width or hd[2],height,text,SZ_PANEL,INK)


def grouped_bars(slide,groups,ymax,ylab,x0=BAR_X0,w=BAR_W,base=BAR_BASE,top=BAR_TOP,
                 ticks=None,value_fmt="{:.1f}%",cat_size=SZ_CAT,k_gap=None):
    """groups = [(label, [(value, color), ...]), ...]. Rounded-top bars, matching slide 20."""
    span=base-top; scale=span/ymax
    ticks = ticks if ticks is not None else [0,20,40,60,80]
    for tv in ticks:
        y=base-tv*scale
        _rect(slide,x0,y,w,0.01,GRID)
        _tx(slide,x0-0.62,y-0.15,0.52,0.34,str(tv),SZ_AXIS,MUTED,
            align=PP_ALIGN.RIGHT,wrap=False)
    _rect(slide,x0,base-0.02,w,0.02,AXIS)
    # rotated y-axis title: build the box HORIZONTAL then rotate 270 about its centre,
    # so the text lays out on one line instead of wrapping one character per row.
    cx=x0-1.30; cy=(top+base)/2.0
    r=_tx(slide,cx-1.30,cy-0.19,2.60,0.38,ylab,SZ_AXIS,MUTED,
          align=PP_ALIGN.CENTER,wrap=False)
    r.rotation=270
    n=len(groups); gw=w/n
    for gi,(glab,bars) in enumerate(groups):
        gc=x0+gw*(gi+0.5)
        k=len(bars)
        gap=0.22 if k_gap is None else k_gap
        tot=k*BARW+(k-1)*gap
        bx=gc-tot/2
        for val,col in bars:
            h=val*scale
            _rect(slide,bx,base-h,BARW,h,col,MSO_SHAPE.ROUND_2_SAME_RECTANGLE,adj=0.08)
            _tx(slide,bx+BARW/2-0.39,base-h-0.42,0.78,0.34,value_fmt.format(val),
                SZ_VALUE,col,bold=True,align=PP_ALIGN.CENTER,wrap=False)
            bx+=BARW+gap
        _tx(slide,gc-1.30,base+0.17,2.60,0.38,glab,cat_size,INK,align=PP_ALIGN.CENTER)


def forest(slide,rows,xmin,xmax,ticks,axis_label,x0=FOR_X0,x1=FOR_X1,
           top=FOR_TOP,dy=FOR_DY,null=1.0,lbl_x=FOR_LBL,lbl_w=3.80,cat_w=1.60,
           grid_top=3.74,note_text=None):
    """rows = [(label, point, lo, hi, estimate_text, sub_text, color), ...]. Matches slide 20B.

    Lays the axis ticks and axis title BELOW the last row, computed from the row count,
    so a 5- or 6-row forest cannot collide with them."""
    span=x1-x0; sc=span/(xmax-xmin)
    last=top+dy*(len(rows)-1)
    grid_bot=last+0.62
    tick_y=grid_bot+0.06
    for tv in ticks:
        x=x0+(tv-xmin)*sc
        _rect(slide,x,grid_top,0.01,grid_bot-grid_top,GRID)
        _tx(slide,x-0.30,tick_y,0.60,0.33,("%g"%tv),SZ_AXIS,MUTED,
            align=PP_ALIGN.CENTER,wrap=False)
    if xmin<=null<=xmax:
        _rect(slide,x0+(null-xmin)*sc,grid_top,0.02,grid_bot-grid_top,AXIS)
    for i,(lab,pt,lo,hi,est,sub,col) in enumerate(rows):
        y=top+i*dy
        _tx(slide,x0-cat_w-0.22,y+0.02,cat_w,0.35,lab,SZ_CAT,INK,align=PP_ALIGN.RIGHT)
        xl=x0+(max(lo,xmin)-xmin)*sc; xh=x0+(min(hi,xmax)-xmin)*sc
        _rect(slide,xl,y+0.12,max(xh-xl,0.02),0.06,col)
        xp=x0+(pt-xmin)*sc
        _rect(slide,xp-0.125,y+0.03,0.25,0.25,col,MSO_SHAPE.OVAL)
        _tx(slide,lbl_x,y-0.19,lbl_w,0.43,est,SZ_EST,INK,bold=True,wrap=False)
        _tx(slide,lbl_x,y+0.23,lbl_w,0.33,sub,SZ_SUB,MUTED,wrap=False)
    ax_y=tick_y+0.40
    _tx(slide,x0-cat_w-0.22,ax_y,9.60,0.33,axis_label,SZ_AXIS,MUTED,wrap=False)
    if note_text:
        _tx(slide,x0-cat_w-0.22,ax_y+0.38,9.60,0.33,note_text,SZ_AXIS,INK,bold=True,wrap=False)
    return ax_y+0.38


def _readable_on(bg):
    """Pick in-bar text colour by background luminance -- white text on #C5CDD8 is unreadable."""
    lum = 0.299*bg[0] + 0.587*bg[1] + 0.114*bg[2]
    return INK if lum > 150 else WHITE


def stacked_rows(slide,rows,segments,y0=4.06,dy=1.67,label_x=1.00,label_w=3.20):
    """rows = [(line1, line2, sublabel, color, [pct,...]), ...]; segments = [(name,color),...].

    Three-line label block (category / race / n) laid out explicitly so nothing overlaps."""
    for ri,(l1,l2,sub,lcol,vals) in enumerate(rows):
        y=y0+ri*dy
        _tx(slide,label_x,y+0.10,label_w,0.34,l1,SZ_CAT,lcol,bold=True,
            align=PP_ALIGN.RIGHT,wrap=False)
        _tx(slide,label_x,y+0.45,label_w,0.34,l2,SZ_CAT,lcol,bold=True,
            align=PP_ALIGN.RIGHT,wrap=False)
        _tx(slide,label_x,y+0.80,label_w,0.32,sub,SZ_SUB,MUTED,
            align=PP_ALIGN.RIGHT,wrap=False)
        x=STK_X0
        for si,v in enumerate(vals):
            w=STK_W*v/100.0
            col=segments[si][1]
            _rect(slide,x,y,w,STK_H,col)
            # Label INSIDE wherever it fits -- an outside label lands in the inter-row gap
            # and collides with the row above. Colour is chosen by segment luminance.
            if w>0.62:
                _tx(slide,x+w/2-0.45,y+0.41,0.90,0.38,"%.1f%%"%v,
                    SZ_CAT if w>0.95 else SZ_SUB,_readable_on(col),
                    bold=True,align=PP_ALIGN.CENTER,wrap=False)
            x+=w
    return y0+dy*len(rows)


def segment_legend(slide,segments,y=8.43,x0=2.17):
    x=x0
    for name,col in segments:
        _rect(slide,x,y+0.07,0.24,0.18,col)
        w=0.112*len(name)+0.20
        _tx(slide,x+0.33,y,w,0.32,name,SZ_SUB,INK)
        x+=0.33+w+0.42


def note(slide,text,l=10.43,t=8.38,w=7.79,h=0.33,color=INK,size=SZ_AXIS,bold=False):
    _tx(slide,l,t,w,h,text,size,color,bold=bold)


def blank(prs, layout=None):
    """New slide on the deck's own layout, with any placeholders removed."""
    lay = layout if layout is not None else (
        prs.slides[19].slide_layout if len(prs.slides) > 19 else prs.slide_layouts[0])
    s=prs.slides.add_slide(lay)
    for ph in list(s.shapes):
        if ph.is_placeholder:
            ph._element.getparent().remove(ph._element)
    return s


def move_slide(prs,old_index,new_index):
    """Reorder by manipulating sldIdLst (python-pptx has no API for this)."""
    lst=prs.slides._sldIdLst
    ids=list(lst)
    lst.remove(ids[old_index])
    lst.insert(new_index,ids[old_index])


# ---- log-symmetric ratio axis ------------------------------------------------
# Exact reciprocal pairs, so the null sits dead centre and the tick labels mirror.
_LADDER = [(1.25, 1.6), (1.25, 1.5), (1.6, 2.5), (1.4, 2.0), (2.0, 4.0)]

def sym_axis(values):
    """Pick [1/k, k] and five mirror-symmetric ticks covering every value/CI bound."""
    import math
    ext = max(max(values), 1.0/min(values))
    for m, k in sorted(_LADDER, key=lambda t: t[1]):
        if ext <= k * 1.001:
            return 1.0/k, k, [round(1.0/k, 3), round(1.0/m, 3), 1.0, m, k]
    k = 4.0; m = 2.0
    return 1.0/k, k, [1.0/k, 1.0/m, 1.0, m, k]


def forest_log(slide, rows, axis_label, x0, x1, top, dy,
               lbl_x, lbl_w=3.80, cat_w=3.20, grid_top=None, note_text=None,
               tick_fmt=lambda v: ("%g" % round(v, 2))):
    """Forest on a log ratio scale with the null (1.0) centred.

    A linear ratio axis compresses protective effects and stretches harmful ones, so 0.5 and 2.0
    -- the same effect in opposite directions -- sit at different distances from the null. On a
    log axis they are equidistant, which is what makes a centred reference line meaningful.
    """
    import math
    vals=[]
    for _, pt, lo, hi, *_ in rows:
        vals += [pt, lo, hi]
    lo_x, hi_x, ticks = sym_axis(vals)
    L0, L1 = math.log(lo_x), math.log(hi_x)
    px = lambda v: x0 + (math.log(v) - L0) / (L1 - L0) * (x1 - x0)

    last = top + dy * (len(rows) - 1)
    gtop = grid_top if grid_top is not None else top - 0.55
    gbot = last + 0.62
    tick_y = gbot + 0.06
    for tv in ticks:
        x = px(tv)
        is_null = abs(tv - 1.0) < 1e-9
        _rect(slide, x, gtop, 0.02 if is_null else 0.01, gbot - gtop, AXIS if is_null else GRID)
        _tx(slide, x - 0.34, tick_y, 0.68, 0.33, tick_fmt(tv), SZ_AXIS,
            INK if is_null else MUTED, bold=is_null, align=PP_ALIGN.CENTER, wrap=False)

    for i, (lab, pt, lo, hi, est, sub, col) in enumerate(rows):
        y = top + i * dy
        _tx(slide, x0 - cat_w - 0.22, y + 0.02, cat_w, 0.35, lab, SZ_CAT, INK,
            align=PP_ALIGN.RIGHT, wrap=False)
        xl, xh = px(max(lo, lo_x)), px(min(hi, hi_x))
        _rect(slide, xl, y + 0.12, max(xh - xl, 0.02), 0.06, col)
        xp = px(pt)
        _rect(slide, xp - 0.125, y + 0.03, 0.25, 0.25, col, MSO_SHAPE.OVAL)
        _tx(slide, lbl_x, y - 0.19, lbl_w, 0.43, est, SZ_EST, INK, bold=True, wrap=False)
        _tx(slide, lbl_x, y + 0.23, lbl_w, 0.33, sub, SZ_SUB, MUTED, wrap=False)

    ax_y = tick_y + 0.42
    lx = x0 - cat_w - 0.22
    bw = min(9.6, 19.06 - lx)          # never run past the deck's right margin
    _tx(slide, lx, ax_y, bw, 0.33, axis_label, SZ_AXIS, MUTED, wrap=False)
    if note_text:
        _tx(slide, lx, ax_y + 0.38, bw, 0.33, note_text, SZ_AXIS, INK, bold=True, wrap=False)
    return ax_y + 0.38


def clear_slide(slide, keep=None):
    """Strip every shape from a slide so it can be redrawn in place.

    Rewriting a slide in place is the ONLY safe way to replace one: deleting a slide with
    part.drop_rel() frees its relationship id, and the next add_slide() reuses that id, which
    silently re-points earlier sldId entries at the new slide. That corrupted this deck once
    (four different slides all rendered as the last one built) -- never delete, always rewrite.
    """
    keep = keep or (lambda sh: False)
    for sh in list(slide.shapes):
        if not keep(sh):
            sh._element.getparent().remove(sh._element)
    return slide


def find_slide(prs, title_prefix):
    """Index of the first slide whose any text box starts with title_prefix, else None."""
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(title_prefix):
                return i
    return None


# =============================================================================
# SAFE DECK OPERATIONS
#
# Everything below was learned the hard way editing a live deck while the author
# edited it in parallel. Read references/pptx-gotchas.md before changing any of it.
# =============================================================================

from pptx.oxml.ns import qn


def find_slide_by_title(prs, title, exact=True):
    """Locate a slide by its title text. NEVER address slides by index.

    Indices move whenever the author adds or deletes a slide, which they do
    between your runs. Every write must re-locate its target by content.
    """
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if (t == title) if exact else (title in t):
                return s
    return None


def keep_furniture(slide, title_rule_y=1.70, footer_rule_y=10.20):
    """Delete the slide body, keep the furniture. Returns shapes removed.

    Furniture is located GEOMETRICALLY -- anything above the title rule or below
    the footer rule -- not by shape index. Index-based retention breaks the first
    time the author edits the slide.
    """
    from pptx.util import Emu
    removed = 0
    for sh in list(slide.shapes):
        top = Emu(sh.top).inches if sh.top is not None else 0.0
        if top <= title_rule_y or top >= footer_rule_y:
            continue
        sh._element.getparent().remove(sh._element)
        removed += 1
    return removed


def insert_slide_after(prs, ref_slide, layout=None):
    """Add a slide and move it directly after ref_slide.

    add_slide() always appends; ordering is fixed by reordering p:sldIdLst.
    NEVER call drop_rel() -- it frees relationship IDs that add_slide() then
    reuses, silently corrupting the deck. That prohibition is absolute.
    """
    idx = list(prs.slides).index(ref_slide)
    new = prs.slides.add_slide(layout or ref_slide.slide_layout)
    for sh in list(new.shapes):              # strip layout placeholders
        sh._element.getparent().remove(sh._element)
    lst = prs.slides._sldIdLst
    moved = list(lst)[-1]
    lst.remove(moved)
    lst.insert(idx + 1, moved)
    return new


def full_bleed_picture(prs, slide, img_path):
    """Place a rendered figure edge to edge. The figure must be 16:9."""
    return slide.shapes.add_picture(str(img_path), 0, 0,
                                    width=prs.slide_width,
                                    height=prs.slide_height)


def set_white_background(slide):
    """Drop a slide-level background override so the layout's white shows."""
    cSld = slide._element.find(qn("p:cSld"))
    bg = cSld.find(qn("p:bg"))
    if bg is not None:
        cSld.remove(bg)
        return True
    return False


# ---- outline colour, read safely -------------------------------------------

def line_hex(shape):
    """Read a shape's outline colour from XML.

    Reading shape.line.color in python-pptx calls _get_or_add_ln(), which INSERTS
    an empty <a:ln><a:solidFill/></a:ln>. PowerPoint renders that as a black
    hairline. One audit pass once put 285 stray outlines across 41 slides this
    way. Always read outlines from the XML instead.
    """
    ln = shape._element.find(".//" + qn("a:ln"))
    if ln is None:
        return None
    clr = ln.find(".//" + qn("a:srgbClr"))
    return clr.get("val") if clr is not None else None


def strip_empty_outlines(prs):
    """Repair <a:ln><a:solidFill/></a:ln> stubs left by reading shape.line.color."""
    fixed = 0
    for s in prs.slides:
        for ln in s._element.iter(qn("a:ln")):
            kids = list(ln)
            if len(kids) == 1 and kids[0].tag == qn("a:solidFill") and not len(list(kids[0])):
                ln.getparent().remove(ln)
                fixed += 1
    return fixed


# ---- native tables ----------------------------------------------------------

_EDGES = ["lnL", "lnR", "lnT", "lnB"]


def cell_borders(cell, spec):
    """spec: {'lnT': (hex_or_None, pt), ...}; omitted edges are cleared.

    The edge elements must appear in schema order (L, R, T, B) and BEFORE the
    fill element, so the whole tcPr child list is rebuilt. Note the tag is
    'a:lnT', not 'a:ln' + 'lnT' -- getting that wrong produces <a:lnlnT>, which
    PowerPoint silently ignores.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    for e in _EDGES:
        for old in tcPr.findall(qn("a:" + e)):
            tcPr.remove(old)
    rest = list(tcPr)
    for ch in rest:
        tcPr.remove(ch)
    for e in _EDGES:
        col, pt = spec.get(e, (None, 1.0))
        ln = tcPr.makeelement(qn("a:" + e), {"w": str(int(pt * 12700)),
                                             "cap": "flat", "cmpd": "sng",
                                             "algn": "ctr"})
        if col is None:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        else:
            fill = ln.makeelement(qn("a:solidFill"), {})
            fill.append(fill.makeelement(qn("a:srgbClr"), {"val": col}))
            ln.append(fill)
        tcPr.append(ln)
    for ch in rest:
        tcPr.append(ch)


def cell_no_fill(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for t in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill"):
        for old in tcPr.findall(qn(t)):
            tcPr.remove(old)
    tcPr.append(tcPr.makeelement(qn("a:noFill"), {}))


def plain_table(slide, rows, cols, l, t, w, h, col_widths=None, row_heights=None):
    """A table with the deck's own styling: no banding, no fills, hairlines only."""
    gf = slide.shapes.add_table(rows, cols, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = gf.table
    tblPr = tbl._tbl.find(qn("a:tblPr"))
    for a in ("firstRow", "bandRow", "firstCol", "bandCol", "lastRow", "lastCol"):
        tblPr.set(a, "0")
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    if row_heights:
        for i, rh in enumerate(row_heights):
            tbl.rows[i].height = Inches(rh)
    for r in range(rows):
        for c in range(cols):
            cell_no_fill(tbl.cell(r, c))
            cell_borders(tbl.cell(r, c), {})
    return tbl


def same_shape(a, b):
    """python-pptx re-wraps shapes on every access, so `a is b` is ALWAYS False
    even for the same shape. Compare the underlying XML elements."""
    return a._element is b._element
