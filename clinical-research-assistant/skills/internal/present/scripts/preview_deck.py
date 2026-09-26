#!/usr/bin/env python3
"""preview_deck.py -- offline layout preview (no PowerPoint, no Keynote).

usage: python3 preview_deck.py DECK.pptx OUTDIR [--dpi 60] [--slides 3,5] [--sheet]

Draws each slide's native shapes (rect / roundRect / ellipse / triangle fills, connectors,
open custGeom polylines) and text (Times New Roman at the real point size, colour, alignment,
anchor, 270-degree rotation) with PIL. It is a LAYOUT check -- collisions, overflow, sizes,
alignment -- not a pixel-exact render. Use it when PowerPoint is busy with the author's own
file (never drive the author's live PowerPoint session) or when app automation is blocked.
The final visual sign-off still comes from a real PowerPoint/Keynote PDF export."""
import argparse, math, os, sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

E = 914400; A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
FONT = "/System/Library/Fonts/Supplemental/Times New Roman.ttf"; FONT_B = "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf"


def hexcol(el, path):
    c = el.find(path)
    return "#" + c.get("val") if c is not None else None


def walk(shapes, off=(0, 0)):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def render(slide, W, H, dpi):
    im = Image.new("RGB", (int(W * dpi), int(H * dpi)), "white"); dr = ImageDraw.Draw(im)
    px = lambda v: v / E * dpi
    for sh in walk(slide.shapes):
        if sh.left is None: continue
        l, t, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        el = sh._element; sp = el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
        fill = hexcol(sp, A + "solidFill/" + A + "srgbClr") if sp is not None else None
        line = hexcol(sp, A + "ln/" + A + "solidFill/" + A + "srgbClr") if sp is not None else None
        g = el.find(".//" + A + "prstGeom"); prst = g.get("prst") if g is not None else None
        cust = el.find(".//" + A + "custGeom")
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            dr.rectangle([l, t, l + w, t + h], outline="#BBBBBB"); continue
        if el.tag.endswith("cxnSp"):
            flipH = el.find(".//" + A + "xfrm").get("flipH") == "1"; flipV = el.find(".//" + A + "xfrm").get("flipV") == "1"
            x1, x2 = (l + w, l) if flipH else (l, l + w); y1, y2 = (t + h, t) if flipV else (t, t + h)
            dr.line([x1, y1, x2, y2], fill=line or "#6B7280", width=max(1, int(dpi / 50))); continue
        if cust is not None:
            path = cust.find(".//" + A + "path"); pw, ph = float(path.get("w") or 1), float(path.get("h") or 1)
            pts = [(l + float(p.get("x")) / pw * w, t + float(p.get("y")) / ph * h) for p in path.iter(A + "pt")]
            closed = path.find(A + "close") is not None
            if closed and fill and len(pts) > 2: dr.polygon(pts, fill=fill)
            elif len(pts) > 1: dr.line(pts, fill=line or "#000000", width=max(1, int(dpi / 30)))
            continue
        if fill:
            if prst == "ellipse": dr.ellipse([l, t, l + w, t + h], fill=fill)
            elif prst in ("roundRect", "round2SameRect"): dr.rounded_rectangle([l, t, l + w, t + h], radius=min(w, h) * 0.12, fill=fill)
            elif prst == "triangle": dr.polygon([(l + w / 2, t + h), (l, t), (l + w, t)] if (sh.rotation or 0) == 180 else [(l + w / 2, t), (l, t + h), (l + w, t + h)], fill=fill)
            else: dr.rectangle([l, t, l + w, t + h], fill=fill)
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame; bp = tf._txBody.find(A + "bodyPr"); anchor = bp.get("anchor") or "t"
            lines = []
            for p in tf.paragraphs:
                runs = [r for r in p.runs if r.text]
                if not runs: lines.append(("", 18, "#000000", False, "l")); continue
                r = runs[0]; pt = r.font.size.pt if r.font.size else 18
                col = "#" + str(r.font.color.rgb) if (r.font.color and r.font.color.type is not None and r.font.color.type == 1) else "#000000"
                al = {"CENTER (2)": "c", "RIGHT (3)": "r"}.get(str(p.alignment), "l")
                lines.append(("".join(x.text for x in runs), pt, col, bool(r.font.bold), al))
            lh = [pt / 72 * dpi * 1.2 for _, pt, *_ in lines]; total = sum(lh)
            rot = sh.rotation in (270.0, 90.0)
            bw, bh = w, h            # python-pptx width/height are the PRE-rotation box
            layer = Image.new("RGBA", (int(max(bw, 1)), int(max(bh, total, 1))), (0, 0, 0, 0)); ld = ImageDraw.Draw(layer)
            y = {"ctr": (bh - total) / 2, "b": bh - total}.get(anchor, 0)
            for (txt, pt, col, bold, al), hh in zip(lines, lh):
                f = ImageFont.truetype(FONT_B if bold else FONT, max(6, int(pt / 72 * dpi)))
                tw = ld.textlength(txt, font=f); x = {"c": (bw - tw) / 2, "r": bw - tw}.get(al, 0)
                ld.text((x, y), txt, font=f, fill=col); y += hh
            if rot:
                layer = layer.rotate(90 if sh.rotation == 270.0 else -90, expand=True)
                cx, cy = l + w / 2, t + h / 2; im.paste(layer, (int(cx - layer.width / 2), int(cy - layer.height / 2)), layer)
            else:
                im.paste(layer, (int(l), int(t)), layer)
    return im


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("deck"); ap.add_argument("out"); ap.add_argument("--dpi", type=int, default=60)
    ap.add_argument("--slides"); ap.add_argument("--sheet", action="store_true"); a = ap.parse_args()
    prs = Presentation(a.deck); W, H = prs.slide_width / E, prs.slide_height / E; os.makedirs(a.out, exist_ok=True)
    want = {int(x) for x in a.slides.split(",")} if a.slides else None; ims = []
    for n, s in enumerate(prs.slides, 1):
        if want and n not in want: continue
        im = render(s, W, H, a.dpi); p = os.path.join(a.out, "p%02d.png" % n); im.save(p); ims.append(im); print(p)
    if a.sheet and ims:
        cols = 2 if len(ims) > 1 else 1; rows = math.ceil(len(ims) / cols); w, h = ims[0].size
        sh = Image.new("RGB", (cols * (w + 8), rows * (h + 8)), (60, 60, 60))
        for i, im in enumerate(ims): sh.paste(im, ((i % cols) * (w + 8), (i // cols) * (h + 8)))
        sp = os.path.join(a.out, "sheet.png"); sh.save(sp); print(sp)


if __name__ == "__main__":
    main()
