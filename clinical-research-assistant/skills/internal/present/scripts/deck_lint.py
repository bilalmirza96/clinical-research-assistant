#!/usr/bin/env python3
"""deck_lint.py -- enforce the CRA house presentation standard (gold standard: ITSOS 2026 deck).

usage:  python3 deck_lint.py DECK.pptx [--minutes 7] [--include-hidden] [--json out.json]
exit 1 if any FAIL. WARN lines need a human decision (usually: confirm, or fix).

Slide checks                                              Notes checks
  S1 font family is Times New Roman           FAIL          N1 banned transition opener           FAIL
  S2 every run >= 28 pt                        FAIL          N2 NCDB called population-based       FAIL
  S3 text colour black / white / navy 1A3255   FAIL          N3 decimal percentage ("30.2%")       WARN
  S4 bold below 40 pt (structural label?)      WARN          N4 em dash                            WARN
  S5 no slide numbers                          FAIL          N5 "receipt of surgery" phrasing      WARN
  S6 title box at (0.938, 0.729)               FAIL          N6 "next" used as a transition > 3x   WARN
  S7 header/footer rules at 1.623 / 10.286     FAIL          N7 notes length vs talk minutes       WARN
  S8 footer legend geometry + 32 pt labels     FAIL          N8 slide has no notes                 WARN
  S9 nothing off the slide                     FAIL
  S10 normAutofit (silently shrinks text)      WARN
  S11 picture on a data slide                  WARN
  S12 em dash in slide text                    WARN
"""
import argparse, json, re, sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

E = 914400
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
OK_TEXT = {"000000", "FFFFFF", "1A3255", "123057"}      # black, white-on-dark, navy structural labels
MIN_PT = 28.0
TITLE_XY = (0.938, 0.729); HDR_Y = 1.623; FTR_Y = 10.286; RULE_X = 0.938; RULE_W = 18.125
LEG = dict(x0=0.92, sw=0.41, sw_top=10.47, txt_top=10.391, pt=32.0)
BANNED = r"^(Furthermore|Moreover|Additionally|Interestingly)\b"


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def fill_hex(sh):
    sp = sh._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
    c = sp.find(A + "solidFill/" + A + "srgbClr") if sp is not None else None   # read XML: .fill getters are safe, .line getters are not
    return c.get("val") if c is not None else None


def text_width_in(text, pt, face="Times New Roman"):
    try:
        from fontTools.ttLib import TTFont
        f = _FONT_CACHE.setdefault(face, TTFont("/System/Library/Fonts/Supplemental/%s.ttf" % face))
        cmap = f.getBestCmap(); hm = f["hmtx"]; upm = f["head"].unitsPerEm
        return sum(hm[cmap.get(ord(c), cmap[ord("n")])][0] for c in text) / upm * pt / 72
    except Exception:
        return 0.5 * len(text) * pt / 72
_FONT_CACHE = {}


def text_extent(sh):
    """Horizontal extent (in) the TEXT actually occupies, not the box: a 19.94 in title box that
    overhangs the slide edge is harmless when its left-aligned text stops at x = 14."""
    l, w = sh.left / E, sh.width / E
    widths = []
    for p in sh.text_frame.paragraphs:
        pt = max([r.font.size.pt for r in p.runs if r.font.size] or [28.0])
        widths.append(text_width_in("".join(r.text for r in p.runs), pt))
    rw = min(w, max(widths or [0.0]))
    al = str(sh.text_frame.paragraphs[0].alignment or "LEFT")
    if "CENTER" in al:
        return l + w / 2 - rw / 2, l + w / 2 + rw / 2
    if "RIGHT" in al:
        return l + w - rw, l + w
    return l, l + rw


def on_card(sh, shapes):
    """True when a text box sits on a filled card (step box, table header cell): bold allowed there."""
    cx, cy = (sh.left + sh.width / 2) / E, (sh.top + sh.height / 2) / E
    for c in shapes:
        if c is sh or c.width is None or geom(c) not in ("roundRect", "rect"):
            continue
        f = fill_hex(c)
        if not f or f in ("FFFFFF", "E9ECF0", "E1E6ED"):
            continue
        if c.left / E <= cx <= (c.left + c.width) / E and c.top / E <= cy <= (c.top + c.height) / E and c.width / E > 3:
            return True
    return False


def geom(sh):
    g = sh._element.find(".//" + A + "prstGeom")
    return g.get("prst") if g is not None else None


def lint(path, minutes=None, include_hidden=False):
    prs = Presentation(path)
    res = []   # (level, slide, code, message)
    add = lambda lvl, n, code, msg: res.append((lvl, n, code, msg))
    notes_words = 0; next_count = 0
    sw_w, sw_h = prs.slide_width / E, prs.slide_height / E
    for n, s in enumerate(prs.slides, 1):
        hidden = s._element.get("show") == "0"
        if hidden and not include_hidden:
            continue
        shapes = list(walk(s.shapes))
        texts = [sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        title_box = next((sh for sh in texts if any(r.font.size and r.font.size.pt >= 40 for p in sh.text_frame.paragraphs for r in p.runs)
                          and sh.top is not None and sh.top / E < 1.5), None)
        rules = [sh for sh in shapes if fill_hex(sh) == "E1E6ED" and sh.width and sh.width / E > 15]
        content_slide = bool(rules)
        # ---- runs
        for sh in texts:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    f = r.font; t = r.text.strip()[:40]
                    if f.name and f.name != "Times New Roman":
                        add("FAIL", n, "S1", f"font {f.name!r}: {t!r}")
                    if f.size is not None and f.size.pt < MIN_PT:
                        add("FAIL", n, "S2", f"{f.size.pt:g} pt < {MIN_PT:g}: {t!r}")
                    col = None
                    try:
                        if f.color is not None and f.color.type is not None:
                            col = str(f.color.rgb)
                    except Exception:
                        col = "theme"
                    if col not in (None, "theme") and col not in OK_TEXT:
                        add("FAIL", n, "S3", f"text colour {col} (grey/colour text): {t!r}")
                    if f.bold and (f.size is None or f.size.pt < 40) and col != "1A3255" and not on_card(sh, shapes):
                        add("WARN", n, "S4", f"bold body text (allowed only for structural labels): {t!r}")
                    if "—" in r.text:
                        add("WARN", n, "S12", f"em dash in slide text: {t!r}")
            bp = sh.text_frame._txBody.find(A + "bodyPr")
            na = bp.find(A + "normAutofit") if bp is not None else None
            if na is not None and na.get("fontScale"):
                add("WARN", n, "S10", f"text shrunk to {int(na.get('fontScale'))/1000:g}% by autofit: {sh.text_frame.text.strip()[:40]!r}")
            txt = sh.text_frame.text.strip()
            if re.fullmatch(r"\d{1,3}", txt) and sh.top / E > sw_h - 1.3 and sh.left / E > sw_w - 3.0:
                add("FAIL", n, "S5", f"slide number {txt!r} at ({sh.left/E:.2f}, {sh.top/E:.2f})")
        # ---- furniture geometry
        if content_slide and title_box is not None:
            tx, ty = title_box.left / E, title_box.top / E
            if abs(tx - TITLE_XY[0]) > 0.03 or abs(ty - TITLE_XY[1]) > 0.03:
                add("FAIL", n, "S6", f"title at ({tx:.3f}, {ty:.3f}); house is {TITLE_XY}")
        for r_ in rules:
            y = r_.top / E; x = r_.left / E; w = r_.width / E
            target = HDR_Y if y < 5 else FTR_Y
            if abs(y - target) > 0.02 or abs(x - RULE_X) > 0.02 or abs(w - RULE_W) > 0.05:
                add("FAIL", n, "S7", f"rule at ({x:.3f}, {y:.3f}) w {w:.3f}; house is ({RULE_X}, {target}) w {RULE_W}")
        if content_slide and len(rules) < 2:
            add("WARN", n, "S7", f"{len(rules)} of 2 rules present (header + footer)")
        # ---- footer legend (swatches = small rounded squares in the footer band)
        sw = sorted([sh for sh in shapes if geom(sh) == "roundRect" and sh.width and sh.width / E < 0.6
                     and sh.top / E > 10.0 and fill_hex(sh)], key=lambda x: x.left)
        if sw:
            if abs(sw[0].left / E - LEG["x0"]) > 0.03:
                add("FAIL", n, "S8", f"legend starts at x {sw[0].left/E:.3f}; house is {LEG['x0']}")
            for s_ in sw:
                if abs(s_.width / E - LEG["sw"]) > 0.02 or abs(s_.height / E - LEG["sw"]) > 0.02 or abs(s_.top / E - LEG["sw_top"]) > 0.03:
                    add("FAIL", n, "S8", f"legend swatch {s_.width/E:.2f}x{s_.height/E:.2f} at y {s_.top/E:.3f}; house is 0.41 sq at {LEG['sw_top']}")
            labels = [sh for sh in texts if sh.top / E > 10.2 and any(abs(sh.left / E - (s_.left + s_.width) / E) < 0.30 for s_ in sw)]
            for sh in labels:
                if any(r.font.size and abs(r.font.size.pt - LEG["pt"]) > 0.1 for p in sh.text_frame.paragraphs for r in p.runs):
                    add("FAIL", n, "S8", f"footer legend label not {LEG['pt']:g} pt: {sh.text_frame.text.strip()[:30]!r}")
        # ---- bounds, pictures
        for sh in shapes:
            if sh.left is None or sh.width is None:
                continue
            l, t, w, h = sh.left / E, sh.top / E, sh.width / E, sh.height / E
            if getattr(sh, "rotation", 0) in (90.0, 270.0):
                cx, cy = l + w / 2, t + h / 2; l, t, w, h = cx - h / 2, cy - w / 2, h, w
            full_bleed = w * h > 0.9 * sw_w * sw_h
            if sh.has_text_frame and sh.text_frame.text.strip() and getattr(sh, "rotation", 0) not in (90.0, 270.0):
                l, r_ = text_extent(sh); w = r_ - l
            if (l < -0.02 or t < -0.02 or l + w > sw_w + 0.02 or t + h > sw_h + 0.02) and not full_bleed \
                    and sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
                what = repr(sh.text_frame.text.strip()[:40]) if sh.has_text_frame else sh.name
                add("FAIL", n, "S9", f"{what} extends off the slide ({l:.2f},{t:.2f},{l+w:.2f},{t+h:.2f})")
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and content_slide:
                add("WARN", n, "S11", f"picture on a data slide ({sh.name}); data must be native shapes")
        # ---- notes
        note = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            note = s.notes_slide.notes_text_frame.text.strip()
        if not note and texts:
            add("WARN", n, "N8", "no presenter notes")
        if note:
            notes_words += len(note.split())
            for para in [q.strip() for q in note.split("\n") if q.strip()]:
                if re.match(BANNED, para):
                    add("FAIL", n, "N1", f"banned transition opener: {para[:50]!r}")
            if re.search(r"NCDB", note) and re.search(r"population[- ](based|database|data)", note, re.I):
                add("FAIL", n, "N2", "NCDB described as population-based (it is hospital-based)")
            for m in re.findall(r"\d+\.\d+\s?%", note):
                add("WARN", n, "N3", f"decimal percentage in notes {m!r}: say it rounded")
            if "—" in note:
                add("WARN", n, "N4", "em dash in notes")
            if re.search(r"receipt of (curative-intent )?surgery", note, re.I):
                add("WARN", n, "N5", "'receipt of surgery' phrasing: prefer 'the difference in surgery' / 'underwent surgery'")
            next_count += len(re.findall(r"\bnext\b", note, re.I))
    if next_count > 3:
        add("WARN", 0, "N6", f"'next' used {next_count}x across notes: vary the transitions")
    if minutes:
        budget = 130 * minutes
        if notes_words > budget * 1.15:
            add("WARN", 0, "N7", f"notes total {notes_words} words, ~{notes_words/130:.1f} min at 130 wpm (> {minutes} min)")
        else:
            add("INFO", 0, "N7", f"notes total {notes_words} words, ~{notes_words/130:.1f} min at 130 wpm")
    return res


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck"); ap.add_argument("--minutes", type=float); ap.add_argument("--include-hidden", action="store_true")
    ap.add_argument("--json")
    a = ap.parse_args()
    res = lint(a.deck, a.minutes, a.include_hidden)
    for lvl, n, code, msg in sorted(res, key=lambda r: ({"FAIL": 0, "WARN": 1, "INFO": 2}[r[0]], r[1], r[2])):
        print(f"{lvl:4}  slide {n:>2}  {code:<4} {msg}" if n else f"{lvl:4}  deck      {code:<4} {msg}")
    nf = sum(1 for r in res if r[0] == "FAIL"); nw = sum(1 for r in res if r[0] == "WARN")
    print(f"\n{'FAIL' if nf else 'PASS'}: {nf} fail, {nw} warn")
    if a.json:
        json.dump([dict(level=l, slide=n, code=c, message=m) for l, n, c, m in res], open(a.json, "w"), indent=1)
    sys.exit(1 if nf else 0)


if __name__ == "__main__":
    main()
