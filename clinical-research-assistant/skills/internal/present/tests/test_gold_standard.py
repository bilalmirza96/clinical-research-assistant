#!/usr/bin/env python3
"""RED/GREEN test for the present module (run: python3 tests/test_gold_standard.py [outdir]).

RED   a deck in the pre-gold style (18 pt grey axis text, bold coloured values, slide number,
      small flat legend) MUST fail deck_lint.
GREEN the same content built with deck_style (gold standard v2) MUST pass deck_lint with zero FAIL.
All numbers are DEMO values -- this file never feeds a real deck."""
import os, sys, subprocess, tempfile
HERE = os.path.dirname(os.path.abspath(__file__)); SCR = os.path.join(HERE, "..", "scripts")
sys.path.insert(0, SCR)
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import deck_style as D

out = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp()
os.makedirs(out, exist_ok=True)


def new_deck():
    prs = Presentation(); prs.slide_width = Inches(20); prs.slide_height = Inches(11.25)
    return prs


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    return s


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


# ---------------- GREEN: gold-standard build -----------------------------------------
g = new_deck()
s = blank(g); D.frame(s, "Receipt of Surgery by Group, Stage I–III", legend=D.RACE_LEGEND)
D.grouped_bars(s, [("Registry A", [(51.0, D.NHW), (23.4, D.NHB)]), ("Registry B", [(46.8, D.NHW), (22.1, D.NHB)])],
               75, "Underwent surgery (%)", ticks=[0, 25, 50, 75], p_labels=["P<.001", "P<.001"])
D.forest(s, [("Registry A matched", 0.57, 0.53, 0.62, "0.57 (0.53–0.62)", "P<.001"),
             ("Registry B matched", 0.53, 0.45, 0.62, "0.53 (0.45–0.62)", "P<.001")],
         0.4, 1.0, [0.4, 0.6, 0.8, 1.0], header="Odds ratio, Group B vs Group A", arrow_text="Less likely in Group B")
notes(s, "The largest difference in treatment was in surgery. About half of one group and a quarter of the other underwent resection.")
s = blank(g); D.frame(s, "Documented Reason for Non-Receipt of Surgery")
D.stacked_rows(s, [("Group A", [77.4, 11.8, 5.0, 5.8]), ("Group B", [83.5, 9.8, 3.0, 3.7])],
               list(zip(["Not planned", "Contraindicated", "Refused", "Died / unknown"], D.PAL_BLUES)))
D.segment_legend(s, list(zip(["Not planned", "Contraindicated", "Refused", "Died / unknown"], D.PAL_BLUES)))
D.note(s, "Chi-square P<.001", l=13.9, t=1.75, w=5.1, align=D.PP_ALIGN.RIGHT)
notes(s, "To understand why these patients did not undergo surgery, we examined the reasons recorded in the registry.")
s = blank(g); D.frame(s, "Analytic Approach")
D.steps(s, [("Unadjusted", "The difference as it appears in the data"), ("Clinical factors", "Age, sex, stage, histology, comorbidity"),
            ("Socioeconomic factors", "Adds insurance, income and education"), ("Propensity matched", "Each patient paired 1:1")])
notes(s, "Each comparison was estimated in four steps.")
s = blank(g); D.frame(s, "Conclusions")
D.numbered_list(s, ["Survival has improved, but the gains have not been shared equally",
                    "One group was about half as likely to undergo surgery at every stage",
                    "About half of the survival gap is mediated through the difference in surgery"])
notes(s, "Our conclusions are threefold.")
green = os.path.join(out, "green.pptx"); g.save(green)

# ---------------- RED: pre-gold habits ------------------------------------------------
r = new_deck(); s = blank(r)
tb = s.shapes.add_textbox(Inches(0.94), Inches(0.73), Inches(18), Inches(0.7)); run = tb.text_frame.paragraphs[0].add_run()
run.text = "Receipt of Surgery"; run.font.size = Pt(43.5); run.font.bold = True; run.font.name = "Times New Roman"
for txt, size, col, bold, x, y in [("Underwent surgery (%)", 18, "4B5563", False, 1, 5), ("51.0%", 21, "123057", True, 3, 4),
                                   ("14", 16.5, "5B6675", False, 18.83, 10.45), ("Non-Hispanic White", 18, "22333F", False, 1.31, 10.45)]:
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(3), Inches(0.4)); q = b.text_frame.paragraphs[0].add_run()
    q.text = txt; q.font.size = Pt(size); q.font.bold = bold; q.font.name = "Times New Roman"; q.font.color.rgb = RGBColor.from_string(col)
red = os.path.join(out, "red.pptx"); r.save(red)

lint = os.path.join(SCR, "deck_lint.py")
rc_red = subprocess.run([sys.executable, lint, red], capture_output=True, text=True)
rc_green = subprocess.run([sys.executable, lint, green], capture_output=True, text=True)
print("RED   exit", rc_red.returncode, "|", rc_red.stdout.strip().splitlines()[-1])
print("GREEN exit", rc_green.returncode, "|", rc_green.stdout.strip().splitlines()[-1])
if rc_green.returncode: print(rc_green.stdout)
assert rc_red.returncode == 1, "RED deck should fail the linter"
assert rc_green.returncode == 0, "GREEN deck should pass the linter"
print("OK  red fails, green passes ->", out)
