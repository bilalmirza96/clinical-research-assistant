#!/usr/bin/env python3
"""Dependency-free tests for hooks/exit_gates.py (run: python3 tests/test_exit_gates.py).

Mirrors the RED/GREEN runner style of
skills/internal/analyze/tests/test_study_design_tools.py and hooks/tests/test_lessons_recall.py:
no pytest, a `check()` helper that prints PASS/FAIL and records failures, exit code 0 iff
everything passed.

Each case builds a real project directory + a synthetic transcript JSONL under tempfile, then
runs the hook as a real subprocess (python3 exit_gates.py) piping JSON on stdin, exactly as
Claude Code's Stop contract does, so a regression in the transcript-parsing, classification,
subprocess-timeout, or fail-open behavior is caught here rather than only at unit level.
"""
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
HOOK = os.path.join(HERE, "..", "exit_gates.py")

# Imported directly (not just subprocess'd) so extract_bash_redirect_targets can be unit
# tested without building a full transcript/hook round trip.
_spec = importlib.util.spec_from_file_location("exit_gates", HOOK)
exit_gates = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(exit_gates)

FAILS = []


def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    if not cond:
        FAILS.append(msg)


def run_hook(payload, timeout=30.0):
    stdin_text = json.dumps(payload)
    proc = subprocess.run(
        [sys.executable, HOOK],
        input=stdin_text,
        capture_output=True,
        text=True,
        timeout=timeout,
    )
    return proc


def write_transcript(path, tool_uses):
    """tool_uses: list of (cwd, name, input_dict). One assistant message per entry."""
    with open(path, "w", encoding="utf-8") as fh:
        for cwd, name, tool_input in tool_uses:
            entry = {
                "type": "assistant",
                "cwd": cwd,
                "message": {
                    "content": [
                        {"type": "tool_use", "name": name, "input": tool_input},
                    ]
                },
            }
            fh.write(json.dumps(entry) + "\n")
        # a non-JSON line must not break parsing
        fh.write("not valid json at all\n")


def find_reports(project_dir):
    reports = os.path.join(project_dir, "Reports")
    if not os.path.isdir(reports):
        return []
    return [f for f in os.listdir(reports) if f.startswith("exit_gates_")]


# ---------------------------------------------------------------- fixtures

ABSTRACT_WITH_VIOLATIONS = """Introduction: The stake here is clear — we test the gate.

Objective: To determine whether the exit gate blocks on a hard voice failure.

Methods: We wrote a short synthetic abstract with a known violation.

Results: The abstract contains an em dash right here — which is a HARD failure. Furthermore, \
it also contains a banned transition word, which is a second HARD failure by itself, and this \
results section is padded out with extra sentences so that it comfortably remains the largest \
section of the draft, well past twice the length of the methods and conclusions sections that \
surround it here.

Conclusions: The gate should block.

[END OF ABSTRACT BODY]
"""

ABSTRACT_NO_SENTINEL_WITH_VIOLATIONS = """Introduction: The stake here is clear — we test the gate.

Objective: To determine whether the exit gate blocks on a hard voice failure when the \
body-end sentinel is absent, which is exactly when voice_check.py prints its advisory \
"note: body-end sentinel ... not found" line to stdout before its JSON.

Methods: We wrote a short synthetic abstract with a known violation and no sentinel line \
at all, so voice_check counts the whole file and also emits its advisory note first.

Results: The abstract contains an em dash right here — which is a HARD failure. Furthermore, \
it also contains a banned transition word, which is a second HARD failure by itself, and this \
results section is padded out with extra sentences so that it comfortably remains the largest \
section of the draft, well past twice the length of the methods and conclusions sections that \
surround it here.

Conclusions: The gate should still block even though there is no sentinel line.
"""

CLEAN_ABSTRACT = """Introduction: The stake here is clear, and we state it in the first person.

Objective: To determine whether a clean abstract passes the exit gate without triggering a block.

Methods: We wrote a short synthetic abstract with no known violations in it at all.

Results: The abstract avoids every banned transition and every AI-tell phrase, and it avoids \
the em dash character entirely, using a plain comma instead, and this results section is padded \
out with additional sentences describing the (hypothetical) findings in plain declarative prose \
so that it comfortably remains the largest section of the draft, well past twice the length of \
the methods and conclusions sections that surround it here, which is exactly what the house \
voice rubric requires of a structured abstract like this one.

Conclusions: The gate should allow the session to stop.

[END OF ABSTRACT BODY]
"""


def main():
    # ---------------------------------------------------------- unit: redirect-target parsing
    # A quoted redirect target containing spaces must be captured whole, not truncated at the
    # first space, for both `>` (double-quoted) and `tee` (single-quoted).
    targets = exit_gates.extract_bash_redirect_targets(
        'python3 script.py > "/tmp/spaced abstract.md"')
    check(targets == ["/tmp/spaced abstract.md"],
          f"double-quoted redirect target with a space is captured whole (got {targets})")

    targets = exit_gates.extract_bash_redirect_targets(
        "python3 script.py >> '/tmp/another spaced report.md'")
    check(targets == ["/tmp/another spaced report.md"],
          f"single-quoted append-redirect target with a space is captured whole (got {targets})")

    targets = exit_gates.extract_bash_redirect_targets(
        'echo hi | tee "/tmp/spaced tee output.md"')
    check(targets == ["/tmp/spaced tee output.md"],
          f"double-quoted tee target with a space is captured whole (got {targets})")

    targets = exit_gates.extract_bash_redirect_targets(
        "echo hi | tee -a '/tmp/spaced tee append.md'")
    check(targets == ["/tmp/spaced tee append.md"],
          f"single-quoted 'tee -a' target with a space is captured whole (got {targets})")

    # unquoted targets (no spaces) still work as before
    targets = exit_gates.extract_bash_redirect_targets("echo hi > /tmp/plain.md")
    check(targets == ["/tmp/plain.md"], f"unquoted redirect target still works (got {targets})")

    tmp_root = tempfile.mkdtemp(prefix="exit_gates_test_")
    try:
        # ------------------------------------------------------------ case 1
        # an abstract with an em dash and "Furthermore" -> block, report exists
        proj1 = os.path.join(tmp_root, "proj1")
        os.makedirs(os.path.join(proj1, "Reports"), exist_ok=True)
        abstract_path = os.path.join(proj1, "Reports", "abstract_test.md")
        with open(abstract_path, "w", encoding="utf-8") as fh:
            fh.write(ABSTRACT_WITH_VIOLATIONS)
        transcript1 = os.path.join(tmp_root, "transcript1.jsonl")
        write_transcript(transcript1, [
            (proj1, "Write", {"file_path": abstract_path, "content": ABSTRACT_WITH_VIOLATIONS}),
        ])

        proc = run_hook({
            "session_id": "s1",
            "transcript_path": transcript1,
            "stop_hook_active": False,
            "cwd": proj1,
            "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case1 (dirty abstract): hook exits 0")
        out = proc.stdout.strip()
        check(bool(out), "case1: stdout is non-empty (a block decision)")
        if out:
            try:
                decision = json.loads(out)
            except (json.JSONDecodeError, ValueError):
                decision = None
            check(decision is not None and decision.get("decision") == "block",
                  "case1: decision == 'block'")
            check(decision is not None and "abstract_test.md" in decision.get("reason", ""),
                  "case1: reason names the failing file")
            check(decision is not None and "Reports" in decision.get("reason", ""),
                  "case1: reason names the report path")
        reports1 = find_reports(proj1)
        check(len(reports1) == 1, f"case1: exactly one report written under Reports/ (found {reports1})")
        if reports1:
            report_text = open(os.path.join(proj1, "Reports", reports1[0]), encoding="utf-8").read()
            check("voice_check" in report_text, "case1: report mentions voice_check")
            check("HARD FAIL" in report_text, "case1: report records a HARD FAIL")

        # ------------------------------------------------------------ case 1b
        # a dirty abstract with NO [END OF ABSTRACT BODY] sentinel -> voice_check prints its
        # "note: body-end sentinel ... not found" advisory line to stdout BEFORE the JSON.
        # A hard failure here must still BLOCK, not be downgraded to a non-blocking TOOL ERROR
        # by a strict json.loads() choking on the leading advisory text.
        proj1b = os.path.join(tmp_root, "proj1b")
        os.makedirs(os.path.join(proj1b, "Reports"), exist_ok=True)
        no_sentinel_path = os.path.join(proj1b, "Reports", "abstract_no_sentinel.md")
        with open(no_sentinel_path, "w", encoding="utf-8") as fh:
            fh.write(ABSTRACT_NO_SENTINEL_WITH_VIOLATIONS)
        transcript1b = os.path.join(tmp_root, "transcript1b.jsonl")
        write_transcript(transcript1b, [
            (proj1b, "Write", {"file_path": no_sentinel_path,
                               "content": ABSTRACT_NO_SENTINEL_WITH_VIOLATIONS}),
        ])

        proc = run_hook({
            "session_id": "s1b",
            "transcript_path": transcript1b,
            "stop_hook_active": False,
            "cwd": proj1b,
            "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case1b (dirty abstract, no sentinel): hook exits 0")
        out = proc.stdout.strip()
        check(bool(out), "case1b: stdout is non-empty (a block decision, not silently allowed)")
        if out:
            try:
                decision = json.loads(out)
            except (json.JSONDecodeError, ValueError):
                decision = None
            check(decision is not None and decision.get("decision") == "block",
                  "case1b: an em dash + banned transition with no sentinel still BLOCKS "
                  "(voice_check's advisory line must not downgrade this to a TOOL ERROR)")
        reports1b = find_reports(proj1b)
        if reports1b:
            report_text = open(os.path.join(proj1b, "Reports", reports1b[0]), encoding="utf-8").read()
            check("**voice_check** — HARD FAIL" in report_text,
                  "case1b: report records voice_check as a HARD FAIL, not a TOOL ERROR "
                  "(the advisory line must not break the --json parse)")

        # ------------------------------------------------------------ case 2
        # a clean abstract -> allow
        proj2 = os.path.join(tmp_root, "proj2")
        os.makedirs(os.path.join(proj2, "Reports"), exist_ok=True)
        clean_path = os.path.join(proj2, "Reports", "abstract_clean.md")
        with open(clean_path, "w", encoding="utf-8") as fh:
            fh.write(CLEAN_ABSTRACT)
        transcript2 = os.path.join(tmp_root, "transcript2.jsonl")
        write_transcript(transcript2, [
            (proj2, "Write", {"file_path": clean_path, "content": CLEAN_ABSTRACT}),
        ])

        proc = run_hook({
            "session_id": "s2",
            "transcript_path": transcript2,
            "stop_hook_active": False,
            "cwd": proj2,
            "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case2 (clean abstract): hook exits 0")
        check(proc.stdout.strip() == "", "case2: no block decision printed (allow)")
        reports2 = find_reports(proj2)
        check(len(reports2) == 1, f"case2: a report is still written for a clean run (found {reports2})")
        if reports2:
            report_text = open(os.path.join(proj2, "Reports", reports2[0]), encoding="utf-8").read()
            check("No hard failures" in report_text, "case2: report records no hard failures")

        # ------------------------------------------------------------ case 3
        # touched only a .py script -> allow, no report at all
        proj3 = os.path.join(tmp_root, "proj3")
        os.makedirs(proj3, exist_ok=True)
        script_path = os.path.join(proj3, "helper.py")
        with open(script_path, "w", encoding="utf-8") as fh:
            fh.write("print('hello')\n")
        transcript3 = os.path.join(tmp_root, "transcript3.jsonl")
        write_transcript(transcript3, [
            (proj3, "Write", {"file_path": script_path, "content": "print('hello')\n"}),
        ])

        proc = run_hook({
            "session_id": "s3",
            "transcript_path": transcript3,
            "stop_hook_active": False,
            "cwd": proj3,
            "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case3 (.py script only): hook exits 0")
        check(proc.stdout.strip() == "", "case3: no block decision printed (allow)")
        check(not os.path.isdir(os.path.join(proj3, "Reports")),
              "case3: no Reports/ directory created, no report written")

        # ------------------------------------------------------------ case 4
        # stop_hook_active True -> allow immediately, even with a would-block transcript
        proc = run_hook({
            "session_id": "s4",
            "transcript_path": transcript1,  # reuse the dirty-abstract transcript from case 1
            "stop_hook_active": True,
            "cwd": proj1,
            "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case4 (stop_hook_active=True): hook exits 0")
        check(proc.stdout.strip() == "", "case4: no block decision printed even though the "
                                          "transcript would otherwise block (loop guard)")

        # ------------------------------------------------------------ case 5
        # missing transcript_path -> allow
        proc = run_hook({
            "session_id": "s5",
            "stop_hook_active": False,
            "cwd": proj1,
            "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case5 (missing transcript_path): hook exits 0")
        check(proc.stdout.strip() == "", "case5: no block decision printed (allow)")

        # a transcript_path that points nowhere should behave the same way
        proc = run_hook({
            "session_id": "s5b",
            "transcript_path": os.path.join(tmp_root, "does-not-exist.jsonl"),
            "stop_hook_active": False,
            "cwd": proj1,
            "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case5b (nonexistent transcript_path): hook exits 0")
        check(proc.stdout.strip() == "", "case5b: no block decision printed (allow)")

        # ------------------------------------------------------------ case 6
        # registry_lint HARD FAIL (an unlabeled key, H1) -> block. registry_lint and
        # house_style were rewired from stdout-marker heuristics onto the shared --json
        # contract (hard/soft/error, exit 0/1/2); this and the next three cases pin that
        # wiring end to end through the real Stop hook, not just the linters' own --json.
        from docx import Document  # already a hard dependency of house_style.py itself

        proj6 = os.path.join(tmp_root, "proj6")
        os.makedirs(os.path.join(proj6, "Reports"), exist_ok=True)
        reg6_path = os.path.join(proj6, "Reports", "MASTER_ANALYSIS_REGISTRY.json")
        reg6_content = json.dumps({
            "NCDB.os.crude_HR": {"label": "", "current": {"value": 1.1, "ci": [1.0, 1.2]}},
        })
        with open(reg6_path, "w", encoding="utf-8") as fh:
            fh.write(reg6_content)
        transcript6 = os.path.join(tmp_root, "transcript6.jsonl")
        write_transcript(transcript6, [
            (proj6, "Write", {"file_path": reg6_path, "content": reg6_content}),
        ])
        proc = run_hook({
            "session_id": "s6", "transcript_path": transcript6, "stop_hook_active": False,
            "cwd": proj6, "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case6 (registry_lint HARD FAIL): hook exits 0")
        out = proc.stdout.strip()
        decision = json.loads(out) if out else None
        check(decision is not None and decision.get("decision") == "block",
              "case6: an unlabeled registry key (H1) BLOCKS the stop")
        reports6 = find_reports(proj6)
        if reports6:
            report_text = open(os.path.join(proj6, "Reports", reports6[0]), encoding="utf-8").read()
            check("**registry_lint** — HARD FAIL" in report_text,
                  "case6: report records registry_lint as HARD FAIL, wired through --json "
                  "(no stdout-marker heuristic)")

        # ------------------------------------------------------------ case 7
        # registry_lint TOOL ERROR (malformed registry JSON) -> does NOT block, reported only.
        proj7 = os.path.join(tmp_root, "proj7")
        os.makedirs(os.path.join(proj7, "Reports"), exist_ok=True)
        reg7_path = os.path.join(proj7, "Reports", "MASTER_ANALYSIS_REGISTRY.json")
        reg7_content = "not valid json{{{"
        with open(reg7_path, "w", encoding="utf-8") as fh:
            fh.write(reg7_content)
        transcript7 = os.path.join(tmp_root, "transcript7.jsonl")
        write_transcript(transcript7, [
            (proj7, "Write", {"file_path": reg7_path, "content": reg7_content}),
        ])
        proc = run_hook({
            "session_id": "s7", "transcript_path": transcript7, "stop_hook_active": False,
            "cwd": proj7, "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case7 (registry_lint TOOL ERROR): hook exits 0")
        check(proc.stdout.strip() == "",
              "case7: malformed registry JSON is a TOOL ERROR, not a hard failure - no block "
              "decision printed")
        reports7 = find_reports(proj7)
        if reports7:
            report_text = open(os.path.join(proj7, "Reports", reports7[0]), encoding="utf-8").read()
            check("**registry_lint** — TOOL ERROR" in report_text,
                  "case7: report records registry_lint as TOOL ERROR (exit 2, error set)")
            check("No hard failures" in report_text,
                  "case7: a TOOL ERROR never counts as a hard failure")

        # ------------------------------------------------------------ case 8
        # house_style HARD FAIL (wrong font) -> block. Also confirms --json never mutates
        # the file it checks (no --check flag is passed by the hook any more).
        proj8 = os.path.join(tmp_root, "proj8")
        os.makedirs(os.path.join(proj8, "Reports"), exist_ok=True)
        docx8_path = os.path.join(proj8, "Reports", "manuscript_test.docx")
        d8 = Document()
        d8.add_paragraph("Hello world").runs[0].font.name = "Arial"
        d8.save(docx8_path)
        transcript8 = os.path.join(tmp_root, "transcript8.jsonl")
        write_transcript(transcript8, [
            (proj8, "Write", {"file_path": docx8_path, "content": "(binary docx)"}),
        ])
        proc = run_hook({
            "session_id": "s8", "transcript_path": transcript8, "stop_hook_active": False,
            "cwd": proj8, "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case8 (house_style HARD FAIL): hook exits 0")
        out = proc.stdout.strip()
        decision = json.loads(out) if out else None
        check(decision is not None and decision.get("decision") == "block",
              "case8: a wrong-font .docx BLOCKS the stop")
        check(Document(docx8_path).paragraphs[0].runs[0].font.name == "Arial",
              "case8: the hook's house_style check never mutated the file (report-only)")
        reports8 = find_reports(proj8)
        if reports8:
            report_text = open(os.path.join(proj8, "Reports", reports8[0]), encoding="utf-8").read()
            check("**house_style** — HARD FAIL" in report_text,
                  "case8: report records house_style as HARD FAIL, wired through --json "
                  "(no 'N violations' stdout-marker heuristic)")

        # ------------------------------------------------------------ case 9
        # house_style TOOL ERROR (a corrupt .docx - not a real zip) -> does NOT block.
        proj9 = os.path.join(tmp_root, "proj9")
        os.makedirs(os.path.join(proj9, "Reports"), exist_ok=True)
        docx9_path = os.path.join(proj9, "Reports", "manuscript_corrupt.docx")
        with open(docx9_path, "w", encoding="utf-8") as fh:
            fh.write("this is not a real docx file, just plain text\n")
        transcript9 = os.path.join(tmp_root, "transcript9.jsonl")
        write_transcript(transcript9, [
            (proj9, "Write", {"file_path": docx9_path, "content": "this is not a real docx file"}),
        ])
        proc = run_hook({
            "session_id": "s9", "transcript_path": transcript9, "stop_hook_active": False,
            "cwd": proj9, "hook_event_name": "Stop",
        })
        check(proc.returncode == 0, "case9 (house_style TOOL ERROR): hook exits 0")
        check(proc.stdout.strip() == "",
              "case9: a corrupt .docx is a TOOL ERROR, not a hard failure - no block decision "
              "printed")
        reports9 = find_reports(proj9)
        if reports9:
            report_text = open(os.path.join(proj9, "Reports", reports9[0]), encoding="utf-8").read()
            check("**house_style** — TOOL ERROR" in report_text,
                  "case9: report records house_style as TOOL ERROR (exit 2, error set)")
            check("No hard failures" in report_text,
                  "case9: a TOOL ERROR never counts as a hard failure")

        # ------------------------------------------------------------ bonus: malformed stdin
        proc = subprocess.run([sys.executable, HOOK], input="{not valid json",
                              capture_output=True, text=True, timeout=10.0)
        check(proc.returncode == 0, "malformed stdin: hook exits 0 (fail open)")
        check(proc.stdout.strip() == "", "malformed stdin: empty stdout")

        proc = subprocess.run([sys.executable, HOOK], input="",
                              capture_output=True, text=True, timeout=10.0)
        check(proc.returncode == 0, "empty stdin: hook exits 0 (fail open)")
        check(proc.stdout.strip() == "", "empty stdin: empty stdout")

    finally:
        shutil.rmtree(tmp_root, ignore_errors=True)

    print()
    print(f"{len(FAILS)} failure(s)")
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main())
